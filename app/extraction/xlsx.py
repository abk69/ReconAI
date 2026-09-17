"""XLSX tabular extraction via openpyxl with sheet/cell provenance."""

from __future__ import annotations

from io import BytesIO
from uuid import UUID

from app.extraction.base import DocumentExtractor
from app.extraction.schemas import ExtractedDocument, ExtractedTable, TableCell, TextBlock


class XLSXExtractor(DocumentExtractor):
    """Extract worksheet tables from XLSX bytes."""

    name = "xlsx_openpyxl"

    def extract(
        self,
        *,
        document_id: UUID,
        data: bytes,
        filename: str,
        mime_type: str,
    ) -> ExtractedDocument:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("openpyxl is required for XLSX extraction.") from exc

        warnings: list[str] = []
        tables: list[ExtractedTable] = []
        blocks: list[TextBlock] = []
        text_parts: list[str] = []

        try:
            workbook = load_workbook(filename=BytesIO(data), data_only=True, read_only=True)
        except Exception as exc:  # noqa: BLE001
            return ExtractedDocument(
                document_id=document_id,
                extractor_name=self.name,
                warnings=[f"Failed to open XLSX: {exc}"],
                metadata={"filename": filename, "mime_type": mime_type},
            )

        try:
            for sheet in workbook.worksheets:
                rows_raw: list[list[str]] = []
                cells: list[TableCell] = []
                for r_idx, row in enumerate(sheet.iter_rows(values_only=False), start=1):
                    values: list[str] = []
                    for c_idx, cell in enumerate(row, start=1):
                        raw = "" if cell.value is None else str(cell.value).strip()
                        values.append(raw)
                        if raw:
                            ref = cell.coordinate
                            cells.append(
                                TableCell(row=r_idx, column=c_idx, value=raw, cell_ref=ref)
                            )
                            blocks.append(
                                TextBlock(
                                    text=raw,
                                    sheet=sheet.title,
                                    cell=ref,
                                    source_type="xlsx_cell",
                                )
                            )
                    if any(values):
                        rows_raw.append(values)

                if not rows_raw:
                    continue

                headers = rows_raw[0]
                data_rows = rows_raw[1:] if len(rows_raw) > 1 else []
                tables.append(
                    ExtractedTable(
                        sheet=sheet.title,
                        headers=headers,
                        rows=data_rows,
                        cells=cells,
                    )
                )
                text_parts.append(sheet.title)
                text_parts.append("\t".join(headers))
                for data_row in data_rows:
                    text_parts.append("\t".join(data_row))
        finally:
            workbook.close()

        full_text = "\n".join(text_parts)
        if not tables:
            warnings.append("XLSX contained no non-empty worksheets.")

        return ExtractedDocument(
            document_id=document_id,
            text_blocks=blocks,
            tables=tables,
            full_text=full_text,
            extractor_name=self.name,
            warnings=warnings,
            metadata={
                "filename": filename,
                "mime_type": mime_type,
                "sheet_count": len(tables),
            },
        )
